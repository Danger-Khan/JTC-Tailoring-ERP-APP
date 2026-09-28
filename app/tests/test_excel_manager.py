"""Regression tests for ExcelManager's on-disk file handling."""
import os
import shutil
import tempfile
import unittest

import openpyxl

from app.excel_manager import ExcelManager

HEADERS = ["id", "item", "category", "quantity", "unit", "reorder_level", "updated"]


class EnsureFileTestCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="excel_mgr_test_")
        self.mgr = ExcelManager(data_dir=self.tmp)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _write_empty_workbook(self, filename):
        """A workbook that exists but holds no header row — the state
        app/data/inventory.xlsx was actually found in."""
        path = os.path.join(self.tmp, filename)
        openpyxl.Workbook().save(path)
        return path

    def test_missing_file_is_created_with_headers(self):
        self.mgr.ensure_file("inventory.xlsx", HEADERS)
        rows = list(openpyxl.load_workbook(
            os.path.join(self.tmp, "inventory.xlsx")).active.iter_rows(values_only=True))
        self.assertEqual(list(rows[0]), HEADERS)

    def test_existing_headerless_file_gets_its_headers_back(self):
        self._write_empty_workbook("inventory.xlsx")
        self.mgr.ensure_file("inventory.xlsx", HEADERS)
        rows = list(openpyxl.load_workbook(
            os.path.join(self.tmp, "inventory.xlsx")).active.iter_rows(values_only=True))
        self.assertEqual(list(rows[0]), HEADERS)

    def test_first_append_survives_an_existing_empty_workbook(self):
        """The bug this guards: ensure_file() used to return early whenever
        the file existed, so an existing-but-empty workbook kept no header
        row. append_record() then dropped its first row into row 1, and the
        next read_records() treated that row as the header — silently losing
        the record."""
        self._write_empty_workbook("inventory.xlsx")
        self.mgr.append_record(
            "inventory.xlsx", HEADERS,
            ["INV-101", "Cotton Lawn", "Fabric", "40", "m", "10", "2026-09-28"])
        records = self.mgr.read_records("inventory.xlsx", HEADERS)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["item"], "Cotton Lawn")
        self.assertEqual(records[0]["quantity"], "40")

    def test_existing_file_with_headers_is_left_alone(self):
        self.mgr.append_record("inventory.xlsx", HEADERS,
                               ["INV-101", "Lining", "Fabric", "5", "m", "2", "2026-09-28"])
        self.mgr.ensure_file("inventory.xlsx", HEADERS)
        records = self.mgr.read_records("inventory.xlsx", HEADERS)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["item"], "Lining")


if __name__ == "__main__":
    unittest.main()
